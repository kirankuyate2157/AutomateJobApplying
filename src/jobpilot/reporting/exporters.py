"""Exporters registered by name: csv, xlsx. Add json/google-sheets the same way.

    @exporters.register("json")
    def export_json(datasets, out_dir) -> list[Path]: ...
"""
from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ..core.registry import exporters
from .datasets import Dataset


@exporters.register("csv")
def export_csv(datasets: dict[str, Dataset], out_dir: Path) -> list[Path]:
    """One CSV per dataset. utf-8-sig so Excel opens non-ASCII text correctly."""
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for name, (cols, rows) in datasets.items():
        path = out_dir / f"{name}.csv"
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:
            writer = csv.writer(fh)
            writer.writerow(cols)
            writer.writerows(rows)
        paths.append(path)
    return paths


@exporters.register("xlsx")
def export_xlsx(datasets: dict[str, Dataset], out_dir: Path) -> list[Path]:
    """One workbook, one sheet per dataset, styled header, frozen top row, filters."""
    out_dir.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for name, (cols, rows) in datasets.items():
        ws = wb.create_sheet(name[:31])                      # Excel sheet names max 31 chars
        ws.append(cols)
        for row in rows:
            ws.append([("" if v is None else v) for v in row])
        for cell in ws[1]:
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), header_fill
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(cols, start=1):
            longest = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:200] if r[i - 1] is not None])
            ws.column_dimensions[get_column_letter(i)].width = min(60, longest + 2)
    stamp = datetime.now().strftime("%Y-%m-%d")
    path = out_dir / f"jobpilot_{stamp}.xlsx"
    wb.save(path)
    return [path]
