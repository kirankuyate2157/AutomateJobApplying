import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from jobpilot.core.registry import exporters
from jobpilot.domain.models import Job
from jobpilot.reporting.datasets import build_datasets
from jobpilot.reporting import exporters as _exp  # noqa: F401  (registers csv/xlsx)
from jobpilot.reporting.report import today_report
from jobpilot.storage.db import Repository


def seeded():
    repo = Repository(":memory:")
    repo.upsert_job(Job("1", "Python Dev", "Acme", "Pune, India", url="http://x/1"), "q", score=80)
    repo.set_job_status("1", "applied")
    repo.add_post("p1", "http://post/1", "Recruiter R", "Hiring python! mail hr@acme.com", "2h", "kw", True, 4, ["Python Developer"])
    repo.add_contact("email", "hr@acme.com", "post", "p1", "mail hr@acme.com")
    repo.record_question("q1", "Favourite colour?", "text", [], None, "", 0, "pending")
    repo.record_search("python dev|India|86400", 10, 1)
    return repo


class ReportingTests(unittest.TestCase):
    def test_contacts_dataset_joins_source(self):
        cols, rows = build_datasets(seeded())["contacts"]
        row = dict(zip(cols, rows[0]))
        self.assertEqual((row["value"], row["from_who"], row["source_url"]), ("hr@acme.com", "Recruiter R", "http://post/1"))

    def test_exports(self):
        repo = seeded()
        with tempfile.TemporaryDirectory() as d:
            csvs = exporters.get("csv")(build_datasets(repo), Path(d) / "csv")
            self.assertTrue(any(p.name == "contacts.csv" for p in csvs))
            (xlsx,) = exporters.get("xlsx")(build_datasets(repo), Path(d) / "x")
            wb = load_workbook(xlsx)
            self.assertIn("contacts", wb.sheetnames)
            self.assertEqual(wb["contacts"]["B2"].value, "hr@acme.com")
            self.assertEqual(wb["applications"]["B2"].value, "Python Dev")

    def test_report_text(self):
        text = today_report(seeded(), 30)
        self.assertIn("applied today          : 1 / 30", text)
        self.assertIn("questions to answer    : 1", text)
        self.assertIn("emails 1", text)


class UiTests(unittest.TestCase):
    def test_answers_roundtrip(self):
        try:
            from jobpilot.ui.app import create_app
        except ImportError:
            self.skipTest("flask not installed")
        from types import SimpleNamespace
        from jobpilot.settings import Settings
        repo = seeded()
        ctx = SimpleNamespace(repo=repo, settings=Settings())
        client = create_app(ctx).test_client()
        self.assertIn("Favourite colour?", client.get("/answers").get_data(as_text=True))
        client.post("/answers/save", data={"key": "q1", "answer": "Blue"})
        self.assertEqual(repo.get_question("q1")["answer"], "Blue")
        self.assertEqual(client.get("/data/contacts").status_code, 200)
        self.assertEqual(client.get("/data/nope").status_code, 404)
        self.assertIn("applied today", client.get("/").get_data(as_text=True))


if __name__ == "__main__":
    unittest.main()
